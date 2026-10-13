"""读取课件（PDF / PPTX），提取文字并缓存到 extracted.txt。

用法：
    python day1_read_pdf.py --extract
"""

import re
import sys
from pathlib import Path

SUPPORTED = (".pptx", ".pdf")
CACHE = Path("extracted.txt")

CHAPTER_PATTERN = re.compile(r"第\s*(\d+)\s*章")


def chapter_key(path: Path) -> tuple:
    """按文件名里的章节号排序，识别不出就排后面。"""
    match = CHAPTER_PATTERN.search(path.name)
    if match:
        return (0, int(match.group(1)))
    return (1, path.name)


def iter_files() -> list:
    """找出所有课件文件，跳过 Office 临时文件，按章节排序。"""
    files = []
    for suffix in SUPPORTED:
        for path in Path(".").glob(f"*{suffix}"):
            # 跳过 ~$xxx.pptx 这类 Office 锁文件
            if path.name.startswith("~$"):
                continue
            files.append(path)
    return sorted(files, key=chapter_key)


def extract_pptx(path: Path) -> list:
    """读取 PPTX，每页幻灯片返回一段文字。"""
    from pptx import Presentation

    prs = Presentation(str(path))
    slides = []
    for slide in prs.slides:
        lines = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs).strip()
                    if text:
                        lines.append(text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        lines.append(" | ".join(cells))
        slides.append("\n".join(lines))
    return slides


def extract_pdf(path: Path) -> list:
    """用 pdfplumber 读取 PDF，每页返回一段文字。"""
    import pdfplumber

    pages = []
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            pages.append(page.extract_text() or "")
    return pages


def main() -> None:
    files = iter_files()
    if not files:
        print("当前文件夹里没找到 PPTX 或 PDF 文件。")
        print("请把课件复制到：", Path(".").resolve())
        return

    all_pages = []

    for path in files:
        suffix = path.suffix.lower()
        try:
            if suffix == ".pptx":
                pages = extract_pptx(path)
            elif suffix == ".pdf":
                pages = extract_pdf(path)
            else:
                continue
        except Exception as exc:
            print(f"{path.name}  →  读取失败，已跳过（{type(exc).__name__}）")
            continue

        chars = sum(len(p) for p in pages)
        print(f"{path.name}  →  {len(pages)} 页，{chars} 字")

        for index, text in enumerate(pages, start=1):
            all_pages.append((path.name, index, text))

    with CACHE.open("w", encoding="utf-8") as handle:
        for name, index, text in all_pages:
            handle.write(f"\n===== {name} | 第 {index} 页 =====\n")
            handle.write(text)
            handle.write("\n")

    total = sum(len(t) for _, _, t in all_pages)
    print("=" * 50)
    print(f"共 {len(files)} 个文件，{len(all_pages)} 页，总字符数：{total}")
    print(f"已保存到：{CACHE.resolve()}")
    print("=" * 50)


if __name__ == "__main__":
    main()
